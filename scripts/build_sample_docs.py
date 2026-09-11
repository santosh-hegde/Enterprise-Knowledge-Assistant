from pathlib import Path

from docx import Document
from fpdf import FPDF

OUT = Path(__file__).resolve().parent.parent / "data"
OUT.mkdir(exist_ok=True)


class DocPDF(FPDF):
    def header(self):
        self.set_font("Helvetica", "B", 10)
        self.cell(0, 8, "Pinnacle Infotech  |  Internal", align="R")
        self.ln(4)

    def footer(self):
        self.set_y(-12)
        self.set_font("Helvetica", "I", 8)
        self.cell(0, 8, f"Page {self.page_no()}", align="C")


def pdf(name: str, title: str, body: str):
    p = DocPDF()
    p.set_auto_page_break(auto=True, margin=18)
    p.add_page()
    p.set_font("Helvetica", "B", 16)
    p.multi_cell(0, 8, title)
    p.ln(4)
    p.set_font("Helvetica", size=11)
    body = (
        body.replace("\u2014", "-")
        .replace("\u2013", "-")
        .replace("\u2018", "'")
        .replace("\u2019", "'")
        .replace("\u201c", '"')
        .replace("\u201d", '"')
    )
    for para in body.strip().split("\n"):
        if para.strip() == "":
            p.ln(3)
            continue
        p.multi_cell(0, 6, para)
        p.ln(1)
    p.output(OUT / name)


def docx(name: str, title: str, body: str):
    d = Document()
    d.add_heading(title, level=1)
    for para in body.strip().split("\n"):
        d.add_paragraph(para)
    d.save(OUT / name)


LEAVE = """
Effective: 1 April 2025
Applies to: full-time employees in India. Probation is the first 6 months.

Leave types
Earned Leave (EL): 18 days each calendar year. Accrues monthly. You cannot take EL during probation except with written HR approval for a family emergency.

Casual Leave (CL): 12 days a year. Meant for short personal work. Max 2 CL in a block unless your manager agrees.

Sick Leave (SL): 12 days a year. For 3 or more calendar days you must upload a doctor's note on the HR portal. During probation you may use SL only; EL and CL start after confirmation.

Carry-forward
Unused Earned Leave can be carried forward to the next calendar year, up to 10 days. The maximum EL balance that may sit in your account at any time is 30 days. Any unused EL above the 10-day carry-forward limit lapses on 31 December. Casual Leave and Sick Leave do not carry forward. They lapse at year end.

This is the detailed rule. The Employee Handbook only summarises it. If the two documents disagree, this Leave Policy wins.

Applying
Apply for EL at least 7 days in advance on the HR portal (Leave > Apply). CL can be applied the same day. Unauthorised absence of 3 days is treated as loss of pay and may start a disciplinary note.

Public holidays
The holiday calendar is published on the intranet each December. Optional holidays: 2 per year from the published optional list.

Maternity and paternity
Maternity leave: 26 weeks, as per the Maternity Benefit Act. Paternity leave: 7 calendar days, to be taken within 30 days of the child's birth. These do not come out of EL.

Compensatory off
If you work a declared holiday with prior manager approval, you get one compensatory off. Use it within 90 days. Comp-offs do not carry forward to the next year.
"""

HANDBOOK = """
Welcome to Pinnacle Infotech. This handbook is a short guide for day-to-day work. Detailed policies live in their own files. Keep this with Leave_Policy.pdf, IT_Policy.pdf, Travel_Policy.docx, Benefits.txt and Code_of_Conduct.pdf.

Work hours
Office hours are 9:30 to 18:30, Monday to Friday, Bengaluru and Hyderabad offices. Core hours are 11:00 to 16:00. Be online during core hours even on WFH days.

Work from home
After probation you may work from home 2 days a week. Monday and Friday WFH in the same week needs manager approval. During probation, WFH is only allowed for a documented medical reason.

Leave (summary)
Full-time staff get 18 days Earned Leave, 12 days Casual Leave and 12 days Sick Leave. Unused EL may be carried forward, up to 10 days; CL and SL lapse. For the full carry-forward cap (30 days max balance) and probation rules, read Leave_Policy.pdf. Do not rely on this summary for payroll questions.

Joining and exit
Notice period is 60 days for engineers and 90 days for people managers. HR can agree a buyout. Return the laptop, access card and ID on the last working day. IT will wipe the device; do not keep copies of source code.

Conduct
Treat colleagues with respect. Harassment reports go to hr-help@pinnaclesoft.example (mailbox is monitored). Gifts from vendors above INR 2,000 must be declared. See Code_of_Conduct.pdf.

Travel
Book official travel through the Corporate Travel desk. International trips need passport, visa, insurance and an IT security briefing. Details are in Travel_Policy.docx.

IT basics
VPN is required off-office. MFA is mandatory. If your laptop is lost, report it within 2 hours. See IT_Policy.pdf.
"""

IT = """
Pinnacle Infotech IT Acceptable Use Policy

Accounts
Every employee gets an email of the form firstname.lastname@pinnaclesoft.example. Do not share your password. Passwords must be at least 12 characters and are rotated every 90 days. MFA (authenticator app) is required for email, VPN and GitHub.

VPN
Remote access uses GlobalProtect. Install it from the IT portal, not from random websites. If you forgot the VPN PIN, reset it on the IT portal (VPN > Reset PIN) or write to it-help@pinnaclesoft.example. Do not send passwords over Slack.

Devices
Company laptops stay encrypted (FileVault / BitLocker). Personal USB drives are not allowed for company data. If a laptop is lost or stolen, call the IT desk and email it-help@pinnaclesoft.example within 2 hours so we can remote-wipe. Delay beyond 2 hours may lead to a security incident record.

Software
Install only from the approved list on the IT portal. Shadow IT (Dropbox, unsanctioned ChatGPT plugins with code pasted in) is not permitted. Need a new tool? Raise a request; IT responds in 3 working days.

Data
Customer data does not go to personal email. Public GitHub for company code is a termination offence. Use the company GitHub org.

Incidents
Phishing: forward the mail to security@pinnaclesoft.example and do not click. If you already clicked, tell IT the same day.

This policy does not cover leave or travel reimbursements.
"""

CONDUCT = """
Pinnacle Infotech Code of Conduct

Respect
We do not tolerate harassment, bullying or discrimination. Report to your manager or to hr-help@pinnaclesoft.example. Retaliation against a reporter is itself a violation.

Conflicts
If a family member works at a vendor or customer, declare it to HR within 7 days of joining or of the relationship starting. You cannot approve invoices for that vendor.

Gifts
Gifts, meals or festival hampers from vendors above INR 2,000 in value must be declared on the HR portal (Ethics > Gifts). Cash gifts are never allowed. If you are unsure of the value, declare it.

Confidentiality
Offer letters, salaries, unpublished product plans and customer lists stay inside the company. Do not post screenshots of internal tools on social media. On LinkedIn you may say you work here; you may not speak as if you represent Pinnacle Infotech without comms approval.

Alcohol and safety
No alcohol in offices. Client dinners: stay within the travel policy meal cap and do not expense alcohol above INR 1,000 per person.

Discipline
Violations can lead to a warning, loss of pay, or termination, depending on severity. HR and Legal decide, not your skip-level manager alone.

Related documents: Employee_Handbook.pdf for working norms, IT_Policy.pdf for systems, Leave_Policy.pdf for time off.
"""

TRAVEL = """
Pinnacle Infotech Travel Policy

Approval
All official travel needs manager approval on the HR portal before you book. Directors approve international trips. Do not buy tickets on a personal credit card and claim later, except in a documented emergency.

Booking
Use the Corporate Travel desk (travel@pinnaclesoft.example). Personal OTAs (MakeMyTrip, personal Expedia) are not reimbursed unless travel desk is closed and your manager mails an exception.

Domestic
Economy airfare. Train 2AC is fine if it is cheaper and the trip is overnight. Hotel cap: INR 6,000 per night in metro cities (Bengaluru, Mumbai, Delhi NCR, Hyderabad, Chennai, Pune, Kolkata) and INR 4,500 elsewhere. Uber/Ola to and from airport is allowed.

International
Before you travel you need:
1. Valid passport (6 months validity)
2. Visa
3. Travel insurance (finance issues a certificate; the group health policy in Benefits.txt does not replace this)
4. Company forex card from finance
5. IT security briefing (VPN check, laptop encryption). IT_Policy.pdf still applies abroad.

International flights are economy. Premium economy is allowed only if the flying time is 8 hours or more and a VP approves in writing.

Bills
Submit expense reports within 7 days of return. Missing boarding passes for flights above INR 10,000 will delay reimbursement. Alcohol: follow Code_of_Conduct.pdf.

Leave during travel
If a trip overlaps a weekend, that is not extra leave. If you take a personal holiday after a client visit, apply EL as usual (Leave_Policy.pdf).
"""

FAQS = """
Pinnacle Infotech - common questions

How do I apply for leave?
HR portal > Leave > Apply. EL needs 7 days' notice. Details: Leave_Policy.pdf.

How many earned leave days can I carry forward?
Up to 10 unused EL days, and your EL balance cannot go above 30. CL and SL do not carry forward. Source: Leave_Policy.pdf. The handbook only has a short version of this.

I am on probation. Can I take vacation?
During the first 6 months you only have sick leave, unless HR approves an emergency EL. Leave_Policy.pdf.

How do I reset my VPN PIN?
IT portal > VPN > Reset PIN, or mail it-help@pinnaclesoft.example. Do not post the PIN on Slack. IT_Policy.pdf.

My laptop was stolen.
Email it-help@pinnaclesoft.example and call the IT desk within 2 hours. IT_Policy.pdf.

What do I need before an international work trip?
Passport, visa, travel insurance, forex card, IT security briefing, and director approval. Travel_Policy.docx and IT_Policy.pdf.

Is there an internet allowance?
INR 1,000 per month after probation. Benefits.txt.

Can I keep unused learning budget for next year?
No. Learning budget lapses. EL carry-forward is a different rule. Benefits.txt and Leave_Policy.pdf.

What is the hotel cap in Bengaluru?
INR 6,000 per night. Travel_Policy.docx.

Who gets gifts from vendors?
Declare anything above INR 2,000. Cash is never ok. Code_of_Conduct.pdf.
"""


if __name__ == "__main__":
    pdf("Leave_Policy.pdf", "Leave Policy", LEAVE)
    pdf("Employee_Handbook.pdf", "Employee Handbook", HANDBOOK)
    pdf("IT_Policy.pdf", "IT Acceptable Use Policy", IT)
    pdf("Code_of_Conduct.pdf", "Code of Conduct", CONDUCT)
    docx("Travel_Policy.docx", "Travel Policy", TRAVEL)
    docx("Company_FAQs.docx", "Company FAQs", FAQS)
    print("wrote sample docs to", OUT)
